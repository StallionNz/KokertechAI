import os

PLUGIN_METADATA = {
    "name": "Generate DOCX",
    "description": "Creates a Word document (.docx) with title and formatted content",
    "version": "1.0.0",
    "tags": ["document", "docx", "word", "generation"],
    "author": "KokertechAI",
    "requires": ["python-docx"],
    "permissions": ['fs_write']
}

COMMAND_NAME = "GENERATE_DOCX"
SCHEMA = {
    "action": "GENERATE_DOCX",
    "filename": "<name of the file, e.g., report.docx>",
    "title": "<document title>",
    "content": "<markdown or plain text content>"
}
WORKSPACE_DIR = r"C:\KokertechAI"

def execute(intent_json):
    filename = intent_json.get("filename", "report.docx")
    title = intent_json.get("title", "Executive Report")
    content = intent_json.get("content", "")

    if not filename.endswith('.docx'):
        filename += '.docx'

    full_path = os.path.abspath(os.path.join(WORKSPACE_DIR, filename.lstrip("\\/")))

    # Security Gate: Prevent path traversal
    if not full_path.startswith(WORKSPACE_DIR):
        return "❌ Security Exception: Target path is outside the workspace."

    try:
        import docx
    except ImportError:
        return "❌ 'python-docx' is not installed. Run: pip install python-docx"

    try:
        doc = docx.Document()
        doc.add_heading(title, 0)
        
        # Basic parsing to handle paragraphs
        for paragraph in content.split('\n\n'):
            if paragraph.strip():
                doc.add_paragraph(paragraph.strip())
        
        doc.save(full_path)
        return f"✅ Successfully generated Word document: {full_path}"
    except Exception as e:
        return f"❌ DOCX generation failed: {str(e)}"